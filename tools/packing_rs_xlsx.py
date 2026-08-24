#!/usr/bin/env python3
"""Export the (c,r,s) packing sweep to an .xlsx workbook, one tab per workload.

Numbers come from packing_rs_sweep.sweep(), i.e. the REAL stage2_core.build_flows(),
so each tab's (8,1,8) row is the shipped traffic table's flow volume.  Nothing is
re-derived here; this file only formats.

Each tab carries two tables:
  1. layer mapping at 128x128, 1-bit cells, 8 bit-planes per weight
  2. the (c,r,s) sweep, with per-c and global minima highlighted

    python3 tools/packing_rs_xlsx.py [out.xlsx]
"""
import math
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, "/home/nizar/noxim3d-dnn/tools")
import packing_rs_sweep as sw
import stage2_core as core

OUT = sys.argv[1] if len(sys.argv) > 1 else "docs/packing_crs_sweep.xlsx"
XB, NB, NODES = sw.XB, sw.NB, 108
TABS = {"ResNet-50 stage-3 bottleneck": "ResNet-50",
        "VGG-16 block 3": "VGG-16",
        "DeiT-S encoder block": "DeiT-S"}

WHITE = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
MUTED = Font(italic=True, color="595959", size=9)
HDR_FILL = PatternFill("solid", fgColor="44546A")
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
BEST_FILL = PatternFill("solid", fgColor="C6EFCE")     # global min
CMIN_FILL = PatternFill("solid", fgColor="E2EFDA")     # min within a c
CUR_FILL = PatternFill("solid", fgColor="FFE699")      # shipped setup
NOFIT_FILL = PatternFill("solid", fgColor="FFC7CE")    # tiles > 108
THIN = Side("thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CTR = Alignment(horizontal="center")

NUM = "#,##0"
PCT = "0.0%"
DEC = "0.000"
DELTA = "+0.0%;-0.0%;0.0%"


def header(ws, row, labels, widths=None):
    for j, lab in enumerate(labels, start=1):
        c = ws.cell(row=row, column=j, value=lab)
        c.font, c.fill, c.border = WHITE, HDR_FILL, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[row].height = 30


def band(ws, row, text, ncols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font, c.fill = BOLD, SUB_FILL
    c.alignment = Alignment(horizontal="left", vertical="center")


def layer_table(ws, row, shapes):
    """Per-layer crossbar mapping + device occupancy.  Returns next free row."""
    band(ws, row, f"Layer mapping   {XB}x{XB} crossbars, 1-bit cells, "
                  f"{NB} bit-planes per INT8 weight", 10)
    row += 1
    header(ws, row, ["Layer", "k", "Cin", "Cout", f"Rows = k^2*Cin",
                     f"Cols = Cout*{NB}", "R_xb", "C_xb", "Crossbars", "Device occ."])
    row += 1
    tb = dev = ncb = 0
    for n, k, ci, co in shapes:
        rr, cc = k * k * ci, co * NB
        R, C = math.ceil(rr / XB), math.ceil(cc / XB)
        occ = (rr / (R * XB)) * (cc / (C * XB))
        for j, v in enumerate([n, k, ci, co, rr, cc, R, C, R * C, occ], start=1):
            cell = ws.cell(row=row, column=j, value=v)
            cell.border = BOX
            if j == 10:
                cell.number_format = PCT
            elif j >= 5:
                cell.number_format = NUM
        tb += rr * cc; dev += R * C * XB * XB; ncb += R * C
        row += 1
    for j, v in enumerate(["TOTAL", None, None, None, None, None, None, None,
                           ncb, tb / dev], start=1):
        cell = ws.cell(row=row, column=j, value=v)
        cell.font, cell.border = BOLD, BOX
        if j == 10:
            cell.number_format = PCT
        elif j == 9:
            cell.number_format = NUM
    row += 1
    c = ws.cell(row=row, column=1,
                value=f"{dev:,} devices hold {tb:,} weight bits. "
                      f"Occupancy is 100% only because every dimension divides {XB}; "
                      f"these blocks were chosen to fit the mesh exactly.")
    c.font = MUTED
    return row + 2


def sweep_table(ws, row, rows):
    """The (c,r,s) sweep with minima highlighted.  Returns next free row."""
    cur = next(d for d in rows if (d["c"], d["r"], d["s"]) == (8, 1, 8))
    fitting = [d for d in rows if d["fits"]]
    gmin = min(fitting, key=lambda d: d["total"])
    cmin = {c: min((d for d in fitting if d["c"] == c), key=lambda d: d["total"])
            for c in (8, 16, 32)}

    band(ws, row, "(c, r, s) packing sweep   "
                  "c = crossbars/tile,  r = input (row) grouping -> reduce traffic,  "
                  "s = output (column) grouping -> scatter traffic", 16)
    row += 1
    top = row
    header(ws, row, ["c", "r", "s", "Tiles", f"Fits {NODES}?", "Alloc. CBs",
                     "CBs used", "Unused CBs", "Unused %", "Scatter (B)",
                     "Reduce (B)", "Add (B)", "Total (B)", "MiB/pass",
                     "vs. current", "Note"],
           widths=[5, 5, 5, 7, 9, 11, 10, 11, 10, 14, 14, 12, 14, 11, 11, 30])
    row += 1
    first = row
    for d in rows:
        notes = []
        fill = None
        if d is cur:
            notes.append("current setup"); fill = CUR_FILL
        if d is gmin:
            notes.append("GLOBAL MIN"); fill = BEST_FILL
        elif d in cmin.values():
            notes.append(f"min @ c={d['c']}")
            fill = fill or CMIN_FILL
        if not d["fits"]:
            notes.append(f"needs {d['tiles']} tiles > {NODES}")
        if d["reduce"] == 0:
            notes.append("no reduce traffic (r spans every row group)")

        vals = [d["c"], d["r"], d["s"], d["tiles"], "Yes" if d["fits"] else "No",
                d["alloc"], d["used"], d["unused"], d["unused"] / d["alloc"],
                d["scatter"], d["reduce"], d["add"], d["total"],
                d["total"] / 1048576, d["total"] / cur["total"] - 1,
                "; ".join(notes)]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=j, value=v)
            cell.border = BOX
            if j in (6, 7, 8, 10, 11, 12, 13):
                cell.number_format = NUM
            elif j == 9:
                cell.number_format = PCT
            elif j == 14:
                cell.number_format = DEC
            elif j == 15:
                cell.number_format = DELTA
            elif j == 5:
                cell.alignment = CTR
                if not d["fits"]:
                    cell.fill = NOFIT_FILL
            elif j == 16:
                cell.font = MUTED
            if fill and j != 5:
                cell.fill = fill
            if d is gmin and j <= 15:
                cell.font = BOLD
        row += 1

    ws.auto_filter.ref = f"A{top}:P{row - 1}"
    ws.freeze_panes = ws.cell(row=first, column=1)
    for text in ["Minima are over FITTING rows only (Tiles <= 108).",
                 "Reduce traffic depends on r alone: (ceil(R_xb/r) - 1) * Cout * HW * PSUM.",
                 "Scatter traffic depends on s alone: ceil(C_xb/s) * k^2*Cin * HW * ACT.",
                 "r > 1 is not reachable by the shipped converter -- "
                 "stage2_core.Layer asserts R >= ceil(rows/XB). These rows come from an "
                 "in-memory patch, so they are a design-space study, not a runnable config."]:
        ws.cell(row=row, column=1, value=text).font = MUTED
        row += 1
    return row


def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for modname, mesh, tag in sw.WORKLOADS:
        _, used, rows = sw.sweep(modname, mesh, tag, quiet=True)
        shapes = [(n, k, ci, co) for (n, k, ci, co, _) in core.LAYERS]
        ws = wb.create_sheet(TABS[tag])

        ws.cell(row=1, column=1, value=tag).font = TITLE
        psum = getattr(core, "BYTES_PER_PSUM", 4)
        act = getattr(core, "BYTES_PER_ACT", 1)
        ws.cell(row=2, column=1,
                value=f"mesh 6x6x3 = {NODES} nodes   |   crossbar {XB}x{XB}   |   "
                      f"crossbars required = {used:,}   |   "
                      f"activations {act*8}-bit, partial sums {psum*8}-bit   |   "
                      f"HW = {core.HW}   |   bytes are per inference pass").font = MUTED

        nxt = layer_table(ws, 4, shapes)
        sweep_table(ws, nxt, rows)
        ws.sheet_view.showGridLines = False

    wb.save(OUT)
    print(f"wrote {OUT}  ({len(wb.sheetnames)} tabs: {', '.join(wb.sheetnames)})")


if __name__ == "__main__":
    main()
