#!/usr/bin/env python3
"""packing_crs_sweep.xlsx + port-load columns (PIL / PEL / PF), one tab per workload.

Same layout, styling and minima-highlighting as tools/packing_rs_xlsx.py -- this
only appends the port-load block and does NOT drop rows whose tile count exceeds
the mesh.  PIL/PEL are placement-invariant (they are per-tile port rates), so they
are well defined even for packings that need more than 108 tiles; those rows stay
in, flagged as before.

New columns
  PIL, PEL     peak injection / ejection port load, flits/cycle (capacity 1.0)
  PF           = max(PIL, PEL), the port floor
  bound        which side binds
  sust PF      TIME-AVERAGED port load (PF is a peak over one phase window)
  burst        PF / sust PF
  k_max        1 / sust PF -- load multiplier at sustained port saturation
  Feasible     tiles <= 108 AND sust PF <= 1   (the two real criteria)

    python3 tools/packing_rs_xlsx_pf.py [out.xlsx]
"""
import collections
import importlib
import math
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, "/home/nizar/noxim3d-dnn/tools")
import packing_rs_sweep as sw
import packing_rs_xlsx as X
import stage2_core as core

OUT = sys.argv[1] if len(sys.argv) > 1 else "docs/packing_crs_sweep_pf.xlsx"

# LOAD_SCALE and CYCLES_PER_MAC as PUBLISHED in each workload's traffic tables.
# stage2_core defaults to 0.05 / 2e-4; ResNet's tables were generated at 0.026 and
# VGG's at CPM 2.5e-5, so the defaults would put PF 1.92x high for ResNet.  pir is
# linear in LOAD_SCALE, so PIL/PEL/PF/sust scale with it (ratios do not).
PUBLISHED = {
    "stage2_resnet":    dict(DNN_LOAD_SCALE="0.026", DNN_CYCLES_PER_MAC="0.0002"),
    "stage2_vgg":       dict(DNN_LOAD_SCALE="0.05",  DNN_CYCLES_PER_MAC="2.5e-05"),
    "stage2_vitsmall":  dict(DNN_LOAD_SCALE="0.05",  DNN_CYCLES_PER_MAC="0.0002"),
}
XB, NB, NODES = sw.XB, sw.NB, 108
F = 16                                    # flits per packet
PFHDR_FILL = PatternFill("solid", fgColor="2F6FB2")   # port-load block header
INF_FILL = PatternFill("solid", fgColor="FFC7CE")     # sustained > 1


def port_loads(trows):
    """(PIL, PEL, sust_PIL, sust_PEL) from stage2 flow rows -- placement-free."""
    if not trows:
        return 0.0, 0.0, 0.0, 0.0
    period = trows[0]["t_period"]
    cuts = sorted({0, period} | {r["t_on"] for r in trows}
                  | {r["t_off"] for r in trows})
    pil = pel = 0.0
    for a, b in zip(cuts, cuts[1:]):
        inj, ejc = collections.Counter(), collections.Counter()
        for r in trows:
            if r["t_on"] <= a and r["t_off"] >= b:
                inj[r["src"]] += r["pir"] * F
                ejc[r["dst"]] += r["pir"] * F
        if inj:
            pil = max(pil, max(inj.values()))
        if ejc:
            pel = max(pel, max(ejc.values()))
    sinj, sejc = collections.Counter(), collections.Counter()
    for r in trows:
        w = r["pir"] * (r["t_off"] - r["t_on"]) * F / period
        sinj[r["src"]] += w
        sejc[r["dst"]] += w
    return (pil, pel,
            max(sinj.values()) if sinj else 0.0,
            max(sejc.values()) if sejc else 0.0)


def sweep_with_pf(modname, mesh, tag):
    """sweep() re-run so the per-(c,r,s) flow rows can be kept for port loads."""
    m = importlib.import_module(modname)
    importlib.reload(m)
    os.environ["DNN_MESH"] = mesh
    for k, v in PUBLISHED.get(modname, {}).items():
        os.environ[k] = v
    importlib.reload(core)
    core.configure(m)
    shapes = [(n, k, ci, co) for (n, k, ci, co, _) in core.LAYERS]
    used = sum(math.ceil(k * k * ci / XB) * math.ceil(co * NB / XB)
               for _, k, ci, co in shapes)

    rows = []
    for c in (8, 16, 32):
        r = 1
        while r <= c:
            s = c // r
            grid, tiles = {}, 0
            for n, k, ci, co in shapes:
                Rx, Cx = math.ceil(k * k * ci / XB), math.ceil(co * NB / XB)
                R, C = math.ceil(Rx / r), max(1, math.ceil(Cx / s))
                grid[n] = (R, C)
                tiles += R * C
            sw.layers_for(m, grid)
            phases, t_period, durations = core.build_phases()
            vol, cls = core.build_flows()
            trows = core.build_rows(vol, cls, phases, t_period)
            agg = collections.Counter()
            for _, d in cls.items():
                for kind, b in d.items():
                    agg[kind] += b
            sca = agg["scatter"] + agg.get("reduce+scatter", 0)
            red, add = agg["reduce"], agg.get("add", 0)
            pil, pel, spil, spel = port_loads(trows)
            pf, spf = max(pil, pel), max(spil, spel)
            alloc = tiles * c
            rows.append(dict(c=c, r=r, s=s, tiles=tiles, fits=tiles <= NODES,
                             alloc=alloc, used=used, unused=alloc - used,
                             scatter=sca, reduce=red, add=add,
                             total=sca + red + add,
                             PIL=pil, PEL=pel, PF=pf,
                             bound="inj" if pil >= pel else "ejc",
                             sust=spf, burst=pf / spf if spf else 0.0,
                             kmax=1.0 / spf if spf else float("inf"),
                             feasible=(tiles <= NODES and spf <= 1.0)))
            r *= 2
    return tag, used, rows


HDRS = ["c", "r", "s", "Tiles", f"Fits {NODES}?", "Alloc. CBs", "CBs used",
        "Unused CBs", "Unused %", "Scatter (B)", "Reduce (B)", "Add (B)",
        "Total (B)", "MiB/pass", "vs. current",
        "PIL", "PEL", "PF", "Bound", "Sust. PF", "Burst PF/sust", "k_max",
        "Feasible", "Note"]
WIDTHS = [5, 5, 5, 7, 9, 11, 10, 11, 10, 14, 14, 12, 14, 11, 11,
          9, 9, 9, 7, 10, 13, 9, 10, 34]
NPF = 15                                   # index where the port-load block starts


def sweep_table(ws, row, rows):
    """Same as packing_rs_xlsx.sweep_table plus the port-load block."""
    cur = next(d for d in rows if (d["c"], d["r"], d["s"]) == (8, 1, 8))
    fitting = [d for d in rows if d["fits"]]
    gmin = min(fitting, key=lambda d: d["total"])
    cmin = {c: min((d for d in fitting if d["c"] == c), key=lambda d: d["total"])
            for c in (8, 16, 32)}
    pfmin = min(fitting, key=lambda d: d["PF"])

    X.band(ws, row, "(c, r, s) packing sweep   "
                    "c = crossbars/tile,  r = input (row) grouping -> reduce "
                    "traffic,  s = output (column) grouping -> scatter traffic   "
                    "|   PORT LOADS: PIL/PEL/PF are placement-invariant, "
                    "flits/cycle against a capacity of 1.0", len(HDRS))
    row += 1
    X.header(ws, row, HDRS, widths=WIDTHS)
    for j in range(NPF + 1, len(HDRS) + 1):        # tint the new block's header
        ws.cell(row=row, column=j).fill = PFHDR_FILL
    row += 1
    first = row

    for d in rows:
        notes, fill = [], None
        if d is cur:
            notes.append("current setup")
            fill = X.CUR_FILL
        if d is gmin:
            notes.append("GLOBAL MIN bytes")
            fill = X.BEST_FILL
        elif d in cmin.values():
            notes.append(f"min bytes @ c={d['c']}")
            fill = fill or X.CMIN_FILL
        if d is pfmin:
            notes.append("MIN PF")
        if not d["fits"]:
            notes.append(f"needs {d['tiles']} tiles > {NODES}")
        if d["reduce"] == 0:
            notes.append("no reduce traffic (r spans every row group)")

        vals = [d["c"], d["r"], d["s"], d["tiles"], "Yes" if d["fits"] else "No",
                d["alloc"], d["used"], d["unused"], d["unused"] / d["alloc"],
                d["scatter"], d["reduce"], d["add"], d["total"],
                d["total"] / 1048576, d["total"] / cur["total"] - 1,
                d["PIL"], d["PEL"], d["PF"], d["bound"], d["sust"],
                d["burst"], d["kmax"], "Yes" if d["feasible"] else "NO",
                "; ".join(notes)]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=j, value=v)
            cell.border = X.BOX
            if j in (9,):
                cell.number_format = X.PCT
            elif j in (6, 7, 8, 10, 11, 12, 13):
                cell.number_format = X.NUM
            elif j == 14:
                cell.number_format = "0.000"
            elif j == 15:
                cell.number_format = X.DELTA
            elif j in (16, 17, 18, 20, 21, 22):
                cell.number_format = "0.0000"
            if fill and j <= NPF:
                cell.fill = fill
            if j == 5 and not d["fits"]:
                cell.fill = X.NOFIT_FILL
            if j == 23 and not d["feasible"]:
                cell.fill = INF_FILL
            if j == 18 and d is pfmin:
                cell.fill = X.BEST_FILL
                cell.font = X.BOLD
            if j in (16, 17, 18, 19, 20, 21, 22, 23):
                cell.alignment = X.CTR
        row += 1

    ws.auto_filter.ref = f"A{first - 1}:{get_column_letter(len(HDRS))}{row - 1}"
    ws.freeze_panes = ws.cell(row=first, column=4)
    return row


def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for modname, mesh, tag in sw.WORKLOADS:
        tag, used, rows = sweep_with_pf(modname, mesh, tag)
        ws = wb.create_sheet(X.TABS.get(tag, tag)[:31])
        c = ws.cell(row=1, column=1, value=tag)
        c.font = X.TITLE
        pfmin = min((d for d in rows if d["fits"]), key=lambda d: d["PF"])
        bmin = min((d for d in rows if d["fits"]), key=lambda d: d["total"])
        ws.cell(row=2, column=1,
                value=(f"mesh 6x6x3 = {NODES} nodes,  {XB}x{XB} crossbars, "
                       f"{NB} bit-planes/weight,  crossbars required = {used}   |   "
                       f"min PF: ({pfmin['c']},{pfmin['r']},{pfmin['s']}) "
                       f"PF={pfmin['PF']:.4f}   |   "
                       f"min bytes: ({bmin['c']},{bmin['r']},{bmin['s']})   |   "
                       f"rows needing >{NODES} tiles are KEPT (flagged in "
                       f"'Fits {NODES}?')   |   LOAD_SCALE={core.LOAD_SCALE}, "
                       f"CYCLES_PER_MAC={core.CYCLES_PER_MAC} "
                       f"(as published)")).font = X.MUTED
        shapes = [(n, k, ci, co) for (n, k, ci, co, _) in core.LAYERS]
        row = X.layer_table(ws, 4, shapes)
        row = sweep_table(ws, row + 1, rows)
        for txt in ("PIL / PEL / PF are peaks over a phase window and are "
                    "PLACEMENT-INVARIANT: they depend only on (c,r,s).",
                    "Feasibility is tiles <= 108 AND sustained port rate <= 1 "
                    "-- NOT PF <= 1.  PF is a peak over a fraction of the period.",
                    "Burst = PF / sustained PF.  k_max = 1 / sustained PF is the "
                    "load multiplier at which the busiest port saturates on average.",
                    "Rows above 108 tiles are retained; their port loads are still "
                    "valid since PIL/PEL are per-tile rates."):
            ws.cell(row=row + 1, column=1, value=txt).font = X.MUTED
            row += 1
        print(f"  {tag}: {len(rows)} packings, "
              f"{sum(1 for d in rows if not d['fits'])} over {NODES} tiles kept")
    wb.save(OUT)
    print(f"\nwrote {OUT}")


if __name__ == "__main__":
    main()
